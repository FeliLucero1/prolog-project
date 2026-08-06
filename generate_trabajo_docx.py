from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUTPUT = Path(__file__).resolve().parent / "Trabajo_Final_Prolog.docx"


def set_global_format(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(8)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph("UNIVERSIDAD / INSTITUCIÓN")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True

    p = doc.add_paragraph("ASIGNATURA: Programación Lógica y Funcional")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True

    doc.add_paragraph("")
    p = doc.add_paragraph("TRABAJO FINAL")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(16)

    p = doc.add_paragraph("Sistema lógico de compra de jugadores de fútbol en Prolog")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(14)
    p.runs[0].bold = True

    doc.add_paragraph("")
    doc.add_paragraph("")

    fields = [
        "Estudiante: [Nombre y Apellido]",
        "Matrícula: [Número de matrícula]",
        "Docente: [Nombre del docente]",
        "Carrera: [Carrera]",
        "Comisión/Grupo: [Grupo]",
        f"Fecha: {date.today().strftime('%d/%m/%Y')}",
    ]
    for row in fields:
        p = doc.add_paragraph(row)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    h.runs[0].font.name = "Times New Roman"
    h.runs[0].font.size = Pt(14 if level == 1 else 12)


def add_paragraphs(doc: Document, paragraphs: list[str]) -> None:
    for text in paragraphs:
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def main() -> None:
    doc = Document()
    set_global_format(doc)
    add_cover(doc)

    add_heading(doc, "1) Título")
    add_paragraphs(
        doc,
        [
            "Sistema lógico de compra de jugadores de fútbol en Prolog: modelado declarativo, "
            "restricciones múltiples y visualización web para soporte de decisiones."
        ],
    )

    add_heading(doc, "2) Resumen (hasta 200 palabras)")
    add_paragraphs(
        doc,
        [
            "Este trabajo presenta el diseño e implementación de un sistema experto en Prolog para apoyar "
            "decisiones de fichajes en fútbol profesional. El modelo representa jugadores, equipos y reglas "
            "de negocio mediante hechos y predicados, incorporando restricciones de presupuesto, edad, cupos "
            "de extranjeros, necesidades por posición y reglas de liga. Además, se integra un componente de "
            "química entre jugadores para mejorar la evaluación de combinaciones de fichaje. La propuesta "
            "incluye una interfaz visual que permite consultar resultados, explorar plantillas y comparar "
            "candidatos de manera intuitiva. Como aporte práctico, se automatiza la adquisición de datos "
            "visuales (logos y fotos) y se agregan mecanismos de robustez frente a errores de red y "
            "desalineación de datos. Los resultados muestran que el enfoque declarativo facilita expresar "
            "criterios complejos con alta trazabilidad y buena mantenibilidad. En síntesis, el proyecto "
            "demuestra la aplicabilidad de la programación lógica para problemas reales de optimización "
            "combinatoria y apoyo a decisiones. Además, incorpora una implementación funcional comparativa en "
            "Racket/Scheme para resolver consultas equivalentes y contrastar estilos de modelado declarativo, "
            "articulando fundamentos teóricos de ambos paradigmas con una implementación usable para "
            "demostración académica [1], [2], [4], [8]."
        ],
    )

    add_heading(doc, "3) Introducción: presentación del tema y contextualización")
    add_paragraphs(
        doc,
        [
            "El mercado de fichajes en fútbol combina variables económicas, deportivas y estratégicas. En la "
            "práctica, un club no decide solo por rendimiento individual; también considera restricciones "
            "presupuestarias, composición de plantilla, nacionalidades, edad, contexto competitivo y "
            "compatibilidad entre jugadores. Este tipo de escenario es apropiado para la programación lógica, "
            "ya que permite expresar reglas y restricciones de forma declarativa.",
            "El problema abordado en este trabajo consiste en modelar estas decisiones en Prolog, de modo que "
            "sea posible consultar qué jugadores puede fichar un equipo, cuáles son los mejores candidatos y "
            "qué combinación maximiza un criterio de rendimiento bajo restricciones. El objetivo no es reemplazar "
            "la decisión humana, sino asistirla con una herramienta explicable y reproducible.",
            "La relevancia académica del tema radica en conectar conceptos centrales de la asignatura —hechos, "
            "reglas, unificación, búsqueda, backtracking, recursión y funciones de orden superior— con un "
            "dominio real y de alto interés. "
            "Desde una perspectiva aplicada, también se trabaja la ingeniería del sistema: separación modular, "
            "documentación, pruebas básicas y una interfaz de usuario que permita comunicar resultados a "
            "personas no técnicas.",
            "El proyecto toma como referencia la documentación oficial de SWI-Prolog [1], literatura clásica "
            "del lenguaje [2], [3], y añade un componente de visualización con Streamlit para mejorar la "
            "experiencia de uso. Asimismo, se incorporan datos de plantillas y material visual desde fuentes "
            "públicas para enriquecer la demostración [5], [6].",
        ],
    )

    add_heading(doc, "4) Marco teórico")
    add_paragraphs(
        doc,
        [
            "La programación lógica es un paradigma declarativo en el que los programas describen relaciones "
            "entre entidades y condiciones de validez, en lugar de detallar paso a paso el procedimiento "
            "algorítmico. Prolog implementa este paradigma sobre cálculo de predicados de primer orden y un "
            "motor de inferencia basado en resolución SLD [2], [3].",
            "En Prolog, una base de conocimiento está compuesta por hechos y reglas. Una consulta intenta probar "
            "si un objetivo se satisface dadas las cláusulas disponibles. La unificación permite enlazar "
            "variables con términos compatibles, mientras que el backtracking explora alternativas cuando una "
            "rama de búsqueda falla [1], [2].",
            "Desde el punto de vista de modelado, las restricciones son predicados adicionales que acotan "
            "soluciones posibles. En problemas de selección de plantillas o fichajes, esto habilita una "
            "representación natural de cupos, límites y condiciones contextuales. Además, la recursión facilita "
            "el procesamiento de listas para cálculos agregados (coste total, rendimiento total, conteos) [3].",
            "En paralelo, el enfoque funcional (Racket/Scheme) permite representar el mismo dominio mediante "
            "estructuras inmutables y pipelines de transformación (`map`, `filter`, `foldl`), con control "
            "explícito de la estrategia de búsqueda. Esta dualidad habilita una comparación didáctica directa "
            "entre paradigmas de la materia [8].",
            "Finalmente, para la capa de interacción, una interfaz web no reemplaza la lógica, sino que actúa "
            "como traductor entre usuario y consultas Prolog. Una buena interfaz debe preservar trazabilidad, "
            "mostrar resultados explicables y manejar errores de forma robusta [7].",
        ],
    )

    add_heading(doc, "5) Desarrollo")
    add_heading(doc, "5.1 Arquitectura del sistema", level=2)
    add_paragraphs(
        doc,
        [
            "El sistema se organiza en tres capas: (a) conocimiento del dominio, (b) reglas de inferencia y "
            "(c) visualización. La capa de conocimiento contiene hechos de jugadores y equipos (posición, edad, "
            "precio, nacionalidad, rendimiento, presupuesto, necesidades). La capa de reglas implementa la lógica "
            "de decisión para consultas como puede_firmar/2, recomendar/2, mejor_fichaje/2 y combinacion_optima/2.",
            "La capa de visualización conecta mediante un puente Python-Prolog y presenta resultados en una app "
            "web. Esta separación mejora mantenibilidad, facilita pruebas y evita mezclar lógica de negocio con "
            "código de interfaz.",
        ],
    )

    add_heading(doc, "5.2 Modelado declarativo y predicados principales", level=2)
    add_paragraphs(
        doc,
        [
            "Se modelan condiciones de fichaje a partir de predicados auxiliares (precio_jugador, edad_jugador, "
            "posicion_jugador, es_extranjero, etc.) y predicados compuestos. Un fichaje válido requiere cumplir "
            "restricciones de presupuesto, posición y edad. Para listas de jugadores se calculan costo total, "
            "conteo de extranjeros y validación de reglas de liga.",
            "El sistema incorpora también una métrica de química entre pares de jugadores y un rendimiento "
            "mejorado por sinergia. Con ello, la combinación óptima deja de ser solo una suma de métricas "
            "individuales y pasa a capturar interacción entre componentes de la plantilla.",
        ],
    )

    add_heading(doc, "5.3 Estrategia de optimización y control de complejidad", level=2)
    add_paragraphs(
        doc,
        [
            "La generación de subconjuntos puede crecer exponencialmente, por lo que se incluyó una versión "
            "optimizada para interfaz: limitar candidatos por rendimiento y tamaño máximo de combinación. "
            "Esta decisión mantiene resultados útiles para demostración sin comprometer tiempos de respuesta.",
            "Se aplicaron además predicados filtrados para excluir jugadores según contexto (por ejemplo, "
            "rivalidad), preservando coherencia entre lógica y visualización.",
        ],
    )

    add_heading(doc, "5.4 Integración visual y datos externos", level=2)
    add_paragraphs(
        doc,
        [
            "La app web incluye paneles de evaluación, recomendaciones, comparación y galerías. Para mejorar la "
            "experiencia, se implementó sincronización local de fotos y logos, de modo que la interfaz no dependa "
            "estrictamente de disponibilidad externa durante la exposición.",
            "Durante el desarrollo surgieron problemas de consistencia entre nombre e imagen al integrar datos "
            "externos. Se resolvió con parseo robusto por fila, indexación por identificador, limpieza de "
            "caché y fallback seguro a avatar local cuando una imagen no es confiable. Esta práctica reduce "
            "riesgo de errores visuales en la entrega final.",
        ],
    )

    add_heading(doc, "5.5 Validación y resultados observados", level=2)
    add_paragraphs(
        doc,
        [
            "La validación incluyó pruebas de compilación, revisión de errores de ejecución y consistencia de "
            "mapeo entre jugadores e imágenes. Se verificó la disponibilidad de plantillas para múltiples clubes "
            "y la existencia de rutas locales para visualización.",
            "Los resultados muestran que el sistema responde correctamente a consultas centrales del dominio y "
            "presenta información de forma comprensible. En especial, la combinación entre razonamiento declarativo "
            "y salida visual mejora la comunicabilidad del proyecto ante público no técnico.",
        ],
    )

    add_heading(doc, "5.6 Implementación comparativa funcional (Racket/Scheme)", level=2)
    add_paragraphs(
        doc,
        [
            "Para cubrir explícitamente el componente funcional de la asignatura, se implementó un módulo en "
            "Racket que resuelve consultas equivalentes al modelo Prolog: viabilidad de fichaje, recomendaciones, "
            "mejor fichaje y combinación óptima acotada. El modelo funcional utiliza datos inmutables (`struct`), "
            "funciones puras y funciones de orden superior.",
            "A diferencia de Prolog, donde la búsqueda y el backtracking son responsabilidad del motor de "
            "inferencia, en Racket la estrategia es explícita: generación recursiva de subconjuntos, filtros "
            "de validez y selección por criterio de rendimiento. Esto permite comparar no solo rendimiento, "
            "sino también claridad semántica, esfuerzo de implementación y control algorítmico.",
            "Como evidencia de calidad, el módulo funcional incorpora pruebas con `rackunit` y una demo de "
            "consola orientada a exposición académica. De este modo, el trabajo deja de ser exclusivamente "
            "lógico y pasa a ser efectivamente comparativo entre paradigmas declarativos.",
        ],
    )

    add_heading(doc, "5.7 Síntesis comparativa entre paradigmas", level=2)
    add_paragraphs(
        doc,
        [
            "En Prolog, la principal ventaja observada es la concisión para expresar restricciones y relaciones: "
            "la pregunta de negocio se traduce en predicados y el motor explora soluciones. En Racket/Scheme, "
            "la ventaja está en el control explícito de transformaciones y composición funcional, favoreciendo "
            "modularidad y trazabilidad del flujo de datos.",
            "La comparación permite concluir que ambos enfoques son válidos para el problema, pero optimizan "
            "objetivos distintos: Prolog reduce complejidad expresiva en búsqueda relacional; Racket/Scheme "
            "aumenta control algorítmico en pipelines determinísticos. Esta complementariedad refleja con "
            "fidelidad los objetivos formativos de Programación Lógica y Funcional.",
        ],
    )

    add_heading(doc, "6) Conclusiones")
    add_paragraphs(
        doc,
        [
            "Este trabajo permitió aplicar de forma práctica los fundamentos de programación lógica en un problema "
            "realista de toma de decisiones. La principal fortaleza del enfoque Prolog fue su capacidad para "
            "expresar restricciones complejas con código legible y trazable, manteniendo coherencia entre reglas "
            "de negocio y resultados de consulta.",
            "En términos personales, el proyecto consolidó una forma de pensar declarativa: primero definir qué "
            "condiciones hacen válida una solución y luego delegar en el motor de inferencia la exploración del "
            "espacio de búsqueda. Esta forma de modelado resultó especialmente útil para dominios con reglas "
            "interdependientes, como el mercado de fichajes.",
            "También fue relevante el aprendizaje de aspectos de ingeniería: separar capas, manejar errores "
            "de integración, trabajar con datos externos y priorizar robustez de la experiencia de usuario. "
            "La etapa de depuración de inconsistencias visuales reforzó la importancia de validar fuentes, "
            "normalizar nombres y diseñar mecanismos de fallback.",
            "Como trabajo futuro se propone: (a) enriquecer la base de conocimiento con métricas avanzadas "
            "(xG, minutos, lesiones), (b) incorporar optimización multiobjetivo y (c) conectar fuentes de datos "
            "actualizadas en tiempo real con control de calidad automatizado.",
        ],
    )

    add_heading(doc, "7) Bibliografía (formato IEEE)")
    refs = [
        "[1] SWI-Prolog, \"SWI-Prolog Reference Manual,\" 2026. [Online]. Available: https://www.swi-prolog.org/pldoc/",
        "[2] I. Bratko, Prolog Programming for Artificial Intelligence, 4th ed. Harlow, UK: Pearson, 2011.",
        "[3] W. F. Clocksin and C. S. Mellish, Programming in Prolog, 5th ed. Berlin, Germany: Springer, 2003.",
        "[4] Learn Prolog Now!, \"An Introduction to Prolog,\" [Online]. Available: http://www.learnprolognow.org/",
        "[5] Transfermarkt, \"FC Barcelona - Club profile,\" 2026. [Online]. Available: https://www.transfermarkt.com/fc-barcelona/startseite/verein/131",
        "[6] EA Sports, \"EA Sports FC Player Ratings,\" 2026. [Online]. Available: https://www.ea.com/games/ea-sports-fc/ratings",
        "[7] Streamlit, \"Streamlit Documentation,\" 2026. [Online]. Available: https://docs.streamlit.io/",
        "[8] Racket Documentation, \"The Racket Guide,\" 2026. [Online]. Available: https://docs.racket-lang.org/guide/",
        "[9] M. Wooldridge, An Introduction to MultiAgent Systems, 2nd ed. Chichester, UK: Wiley, 2009.",
        "[10] S. Russell and P. Norvig, Artificial Intelligence: A Modern Approach, 4th ed. Hoboken, NJ, USA: Pearson, 2021.",
        "[11] FIFA, \"Regulations on the Status and Transfer of Players,\" 2026. [Online]. Available: https://www.fifa.com/legal",
    ]
    for ref in refs:
        p = doc.add_paragraph(ref)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.save(OUTPUT)
    print(f"Documento generado: {OUTPUT}")


if __name__ == "__main__":
    main()
